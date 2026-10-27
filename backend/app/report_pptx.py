"""The report as a PowerPoint deck (16:9), for lab meetings.

Same content as the HTML/PDF report and the same rules: the narrative text is
the report agent's (checked sentence by sentence against the results, or the
template), every number, table and figure is read from the run folder. Each
slide with model-written text says so in its footer.
"""

import re
from datetime import datetime
from pathlib import Path

from . import results
from .report import _fmt

W, H = 13.333, 7.5          # slide size in inches (16:9)
MARGIN = 0.5
UP, DOWN, GREY, DARK = (0xD5, 0x5E, 0x00), (0x00, 0x72, 0xB2), (0x6B, 0x72, 0x80), (0x1F, 0x29, 0x37)


def build(job: dict, job_dir: Path, run_dir: Path, narrative: dict) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    ov = results.overview(run_dir)
    s, p = ov["summary"], ov["resolved_params"]
    sec, src, model = narrative["sections"], narrative["section_sources"], narrative.get("model")
    test, ref = p["test_level"], p["reference_level"]
    cb = (ov["provenance"] or {}).get("code_bank") or {}
    figs = {f["path"]: f for tab in ov["figures"].values() for f in tab}

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    deck = _Deck(prs, model)

    # 1. title
    sl = deck.slide()
    deck.text(sl, f"{test} vs {ref}", 0.8, 2.3, W - 1.6, 1.2, size=44, bold=True)
    deck.text(sl, f"RNA-seq differential expression · {job.get('dataset') or ''} · {s['n_samples']} samples",
              0.8, 3.5, W - 1.6, 0.6, size=20, color=GREY)
    deck.text(sl, f"{datetime.now():%d %B %Y} · run {ov['run_id']} · code bank {cb.get('version', '')}"
                  + (f" (commit {cb['commit'][:7]})" if cb.get("commit") else ""), 0.8, 4.1, W - 1.6, 0.5, size=14, color=GREY)

    # 2. summary with headline numbers
    sl = deck.slide("Summary", src.get("summary"))
    tiles = [(f"{s['n_samples']}", "samples"), (f"{s['n_genes_tested']:,}", "genes tested"),
             (f"{s['n_significant']:,}", "significant"), (f"{s['n_up']:,}", f"higher in {test}"),
             (f"{s['n_down']:,}", f"lower in {test}")]
    tw = (W - 2 * MARGIN - 0.4 * 4) / 5
    for i, (num, lab) in enumerate(tiles):
        x = MARGIN + i * (tw + 0.4)
        color = UP if i == 3 else DOWN if i == 4 else DARK
        deck.text(sl, num, x, 1.3, tw, 0.7, size=32, bold=True, color=color, align="center")
        deck.text(sl, lab, x, 2.0, tw, 0.4, size=13, color=GREY, align="center")
    deck.text(sl, sec["summary"], MARGIN, 2.7, W - 2 * MARGIN, 3.4, size=16)
    deck.text(sl, f"Significant: {s['significance_rule']}. Fold changes are {test} relative to {ref}.",
              MARGIN, 6.3, W - 2 * MARGIN, 0.4, size=11, color=GREY)

    # 3. key findings
    if sec["key_findings"]:
        sl = deck.slide("Key findings", src.get("key_findings"))
        deck.bullets(sl, sec["key_findings"], MARGIN, 1.3, W - 2 * MARGIN, 5.4, size=20)

    # 4. samples
    sl = deck.slide("Sample structure", src.get("sample_structure"))
    deck.image(sl, run_dir, "03_exploratory/pca_plot.png", MARGIN, 1.2, 7.4, 5.6)
    deck.text(sl, sec["sample_structure"], 8.2, 1.4, W - 8.2 - MARGIN, 5.2, size=16)

    # 5. quality control
    qc = ov["qc"]["metrics"]
    flagged = [m for m in qc if m.get("flags")]
    sl = deck.slide("Quality control")
    deck.image(sl, run_dir, "02_qc/library_sizes.png", MARGIN, 1.2, 6.0, 5.6)
    deck.image(sl, run_dir, "02_qc/logcpm_boxplot.png", 6.8, 1.2, 6.0, 5.6)
    if flagged:
        deck.text(sl, "Flagged: " + "; ".join(f"{m['sample_id']} ({m['flags']})" for m in flagged),
                  MARGIN, 6.85, W - 2 * MARGIN, 0.4, size=11, color=UP)

    # 6. volcano / MA
    sl = deck.slide("Differential expression", src.get("gene_highlights"))
    deck.image(sl, run_dir, "05_de_plots/volcano_plot.png", MARGIN, 1.2, 6.2, 4.4)
    deck.image(sl, run_dir, "05_de_plots/ma_plot.png", 6.9, 1.2, 6.0, 4.4)
    deck.text(sl, sec["gene_highlights"], MARGIN, 5.7, W - 2 * MARGIN, 1.3, size=14)

    # 7. top genes
    sl = deck.slide(f"Top genes (by adjusted p-value)")
    for i, (direction, label, color) in enumerate((("up", f"Higher in {test}", UP), ("down", f"Lower in {test}", DOWN))):
        rows = results.de_table(run_dir, direction=direction, sort="padj", order="asc", page_size=10)["rows"]
        x = MARGIN + i * (W / 2 - 0.1)
        deck.text(sl, label, x, 1.2, W / 2 - 0.8, 0.4, size=16, bold=True, color=color)
        deck.table(sl, ["Gene", "log2 FC", "adj. p"],
                   [[r.get("gene_symbol") or r["gene_id"], _fmt(r["log2FoldChange"]), _fmt(r["padj"], "p")] for r in rows]
                   or [["none", "", ""]], x, 1.7, W / 2 - 0.8, 4.8)

    # 8. heatmap
    if (run_dir / "05_de_plots" / "heatmap_top_degs.png").exists():
        sl = deck.slide("Top differentially expressed genes across samples")
        deck.image(sl, run_dir, "05_de_plots/heatmap_top_degs.png", MARGIN, 1.2, W - 2 * MARGIN, 5.8)

    # 9. pathways
    sl = deck.slide("Pathways", src.get("pathway_highlights"))
    deck.text(sl, sec["pathway_highlights"], MARGIN, 1.2, W - 2 * MARGIN, 1.1, size=15)
    en = ov["enrichment"]
    if en and en.get("status") == "done" and en.get("collections"):
        name, c = next(iter(en["collections"].items()))
        gsea = results.enrichment_table(run_dir, c["folder"], "gsea", limit=8)
        if gsea["figure"]:
            deck.image(sl, run_dir, gsea["figure"], MARGIN, 2.4, 6.4, 4.5)
        if gsea["rows"]:
            deck.text(sl, f"GSEA, {name}: positive NES = higher in {test}", 7.1, 2.4, 5.8, 0.4, size=12, color=GREY)
            deck.table(sl, ["Pathway", "NES", "adj. p"],
                       [[_short(r["pathway"]), _fmt(r["NES"]), _fmt(r["padj"], "p")] for r in gsea["rows"]],
                       7.1, 2.85, 5.8, 4.0, size=10)

    # 10. caveats and next steps
    sl = deck.slide("Things to check and next steps", src.get("caveats"))
    deck.text(sl, "Things to check", MARGIN, 1.2, 6.0, 0.4, size=16, bold=True)
    deck.bullets(sl, sec["caveats"] or ["No caveats were found."], MARGIN, 1.7, 6.0, 5.0, size=15)
    deck.text(sl, "Suggested next steps", 6.9, 1.2, 6.0, 0.4, size=16, bold=True)
    deck.bullets(sl, sec["next_steps"], 6.9, 1.7, 6.0, 5.0, size=15)

    # 11. methods
    sl = deck.slide("Methods")
    rows = [["Comparison", f"{test} vs {ref} (column {p['condition']})"],
            ["Design", p.get("design_formula", "")],
            ["Accounted for", ", ".join(p["covariates"]) or "nothing"],
            ["Significance", s["significance_rule"]],
            ["Code bank", f"{cb.get('version', '')} commit {(cb.get('commit') or '')[:7]}"]]
    deck.table(sl, ["Setting", "Value"], rows, MARGIN, 1.2, 5.6, 3.0, size=11)
    deck.text(sl, _plain(ov["methods_md"] or ""), 6.4, 1.2, W - 6.4 - MARGIN, 5.6, size=10)
    n_model = sum(v == "model" for v in src.values())
    deck.text(sl, (f"Narrative text: {n_model} section(s) written by an open-weight model "
                   f"({(model or '').replace('ollama:', '')}) and checked against the results; "
                   if n_model else "Narrative text written from the results by fixed rules; ")
              + "figures, tables and numbers come from the analysis files. Full code: the run's code .zip.",
              MARGIN, 6.6, W - 2 * MARGIN, 0.6, size=10, color=GREY)

    out = job_dir / "report" / "report.pptx"
    prs.save(out)
    return out


class _Deck:
    def __init__(self, prs, model: str | None) -> None:
        self.prs, self.model = prs, model

    def slide(self, title: str | None = None, source: str | None = None):
        sl = self.prs.slides.add_slide(self.prs.slide_layouts[6])  # blank
        if title:
            self.text(sl, title, MARGIN, 0.35, W - 2 * MARGIN, 0.8, size=30, bold=True)
        if source == "model":
            self.text(sl, f"Text by the assistant ({(self.model or '').replace('ollama:', '')}), checked against the results",
                      MARGIN, H - 0.4, W - 2 * MARGIN, 0.3, size=9, color=GREY, align="right")
        return sl

    def text(self, sl, text: str, x, y, w, h, *, size=14, bold=False, color=DARK, align="left"):
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
        box = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        para = tf.paragraphs[0]
        para.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        run = para.add_run()
        run.text = text
        run.font.size, run.font.bold, run.font.color.rgb = Pt(size), bold, RGBColor(*color)
        return box

    def bullets(self, sl, items: list[str], x, y, w, h, *, size=16):
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt
        box = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.space_after = Pt(10)
            run = para.add_run()
            run.text = "•  " + item
            run.font.size, run.font.color.rgb = Pt(size), RGBColor(*DARK)

    def image(self, sl, run_dir: Path, rel: str, x, y, w, h):
        """The figure, as large as fits in the box, keeping its proportions."""
        from PIL import Image
        from pptx.util import Inches
        path = run_dir / rel
        if not path.exists():
            return
        with Image.open(path) as im:
            ratio = im.width / im.height
        bw, bh = (w, w / ratio) if w / ratio <= h else (h * ratio, h)
        sl.shapes.add_picture(str(path), Inches(x + (w - bw) / 2), Inches(y), Inches(bw), Inches(bh))

    def table(self, sl, headers: list[str], rows: list[list], x, y, w, h, *, size=12):
        from pptx.util import Inches, Pt
        n = len(rows) + 1
        shape = sl.shapes.add_table(n, len(headers), Inches(x), Inches(y), Inches(w), Inches(min(h, 0.37 * n)))
        tbl = shape.table
        for j, head in enumerate(headers):
            tbl.cell(0, j).text = head
        for i, r in enumerate(rows, start=1):
            for j, v in enumerate(r):
                tbl.cell(i, j).text = str(v)
        for i in range(n):
            for j in range(len(headers)):
                for para in tbl.cell(i, j).text_frame.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(size)
        tbl.columns[0].width = Inches(w * (0.5 if len(headers) > 2 else 0.35))
        rest = (w - w * (0.5 if len(headers) > 2 else 0.35)) / max(1, len(headers) - 1)
        for j in range(1, len(headers)):
            tbl.columns[j].width = Inches(rest)


def _short(pathway: str, n: int = 48) -> str:
    s = pathway.replace("_", " ")
    return s if len(s) <= n else s[: n - 1] + "…"


def _plain(md: str, limit: int = 1600) -> str:
    """methods.md as plain text for a slide."""
    text = "\n".join(ln for ln in md.splitlines() if not ln.startswith("# "))
    text = re.sub(r"^#+\s*", "", text, flags=re.M)
    text = re.sub(r"\*\*?([^*]+)\*\*?", r"\1", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text if len(text) <= limit else text[:limit].rsplit(" ", 1)[0] + " … (full text in the HTML report)"
  
